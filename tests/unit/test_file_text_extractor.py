"""Unit tests for uploaded file text extraction service."""

from io import BytesIO

import pytest
from docx import Document as DocxDocument

from app.core.services.file_text_extractor import FileTextExtractor
from app.interfaces.external.file_service import ExternalFile
from app.shared.errors.exceptions import FileExtractionError


class TestFileTextExtractor:
    """Tests for binary file extraction and safety limits."""

    def test_extract_text_file_success(self) -> None:
        extractor = FileTextExtractor()
        file = ExternalFile(
            file_id="f1",
            filename="notes.txt",
            content_type="text/plain",
            content="مرحبا بالعالم".encode("utf-8"),
        )

        result = extractor.extract_many([file])

        assert len(result) == 1
        assert result[0].text == "مرحبا بالعالم"
        assert result[0].truncated is False

    def test_extract_docx_file_success(self) -> None:
        extractor = FileTextExtractor()
        doc = DocxDocument()
        doc.add_paragraph("فقرة أولى")
        doc.add_paragraph("فقرة ثانية")
        buffer = BytesIO()
        doc.save(buffer)

        file = ExternalFile(
            file_id="f2",
            filename="contract.docx",
            content_type=(
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            ),
            content=buffer.getvalue(),
        )

        result = extractor.extract_many([file])

        assert len(result) == 1
        assert "فقرة أولى" in result[0].text
        assert "فقرة ثانية" in result[0].text

    def test_extract_pdf_route_uses_pdf_handler(
        self, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        extractor = FileTextExtractor()

        def fake_extract_pdf(_: bytes) -> str:
            return "pdf-text"

        monkeypatch.setattr(extractor, "_extract_pdf_text", fake_extract_pdf)

        file = ExternalFile(
            file_id="f3",
            filename="case.pdf",
            content_type="application/pdf",
            content=b"%PDF-mock",
        )

        result = extractor.extract_many([file])

        assert result[0].text == "pdf-text"

    def test_unsupported_file_type_raises(self) -> None:
        extractor = FileTextExtractor()
        file = ExternalFile(
            file_id="f4",
            filename="image.png",
            content_type="image/png",
            content=b"not-supported",
        )

        with pytest.raises(FileExtractionError) as exc_info:
            extractor.extract_many([file])

        assert "Unsupported file type" in str(exc_info.value)

    def test_truncates_per_file_and_total_budget(self) -> None:
        extractor = FileTextExtractor()
        extractor._max_chars_per_file = 5
        extractor._max_chars_total = 8

        files = [
            ExternalFile(
                file_id="f5",
                filename="a.txt",
                content_type="text/plain",
                content=b"123456789",
            ),
            ExternalFile(
                file_id="f6",
                filename="b.txt",
                content_type="text/plain",
                content=b"abcdef",
            ),
        ]

        result = extractor.extract_many(files)

        assert result[0].text == "12345"
        assert result[0].truncated is True
        assert result[1].text == "abc"
        assert result[1].truncated is True
