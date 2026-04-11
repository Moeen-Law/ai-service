"""
Files Service Generation Adapter

Generates DOCX bytes locally, requests an upload URL from Files service,
uploads binary content, and returns the stored file ID.
"""

from io import BytesIO
import re
from typing import Any, Dict

import httpx
from docx import Document as DocxDocument

from app.infrastructure.config.settings import get_settings
from app.interfaces.external.file_generation_service import (
    FileGenerationRequest,
    FileGenerationServiceInterface,
    GeneratedFile,
)
from app.shared.errors.exceptions import FileGenerationError


class HTTPFileGenerationService(FileGenerationServiceInterface):
    """DOCX generation service backed by Files service upload endpoints."""

    def __init__(self) -> None:
        settings = get_settings()
        self._base_url = settings.FILES_SERVICE_BASE_URL.rstrip("/")
        self._upload_url_path = settings.FILES_SERVICE_UPLOAD_URL_PATH
        self._auth_token = settings.FILES_SERVICE_AUTH_TOKEN
        self._timeout = settings.FILES_SERVICE_TIMEOUT_SECONDS
        self._verify_tls = settings.FILES_SERVICE_VERIFY_TLS
        self._default_bucket = settings.FILES_SERVICE_UPLOAD_BUCKET

    def _build_client(self) -> httpx.AsyncClient:
        return httpx.AsyncClient(
            timeout=self._timeout,
            verify=self._verify_tls,
        )

    def _build_upload_url_endpoint(self) -> str:
        if not self._base_url:
            raise FileGenerationError(
                message="FILES_SERVICE_BASE_URL is not configured",
                details={"field": "FILES_SERVICE_BASE_URL"},
            )

        path = self._upload_url_path
        if not path.startswith("/"):
            path = f"/{path}"
        return f"{self._base_url}{path}"

    def _build_json_headers(self) -> Dict[str, str]:
        headers: Dict[str, str] = {
            "Accept": "application/json",
            "Content-Type": "application/json",
        }
        if self._auth_token:
            headers["Authorization"] = f"Bearer {self._auth_token}"
        return headers

    def _resolve_metadata_value(
        self,
        request: FileGenerationRequest,
        key_variants: list[str],
    ) -> str:
        metadata = request.metadata or {}
        for key in key_variants:
            value = metadata.get(key)
            if isinstance(value, str) and value.strip():
                return value.strip()
        return ""

    def _resolve_bucket(self, request: FileGenerationRequest) -> str:
        bucket = self._resolve_metadata_value(request, ["bucket"])
        if bucket:
            return bucket
        return self._default_bucket

    @staticmethod
    def _sanitize_filename(filename: str) -> str:
        """Normalize filename to Files service accepted characters."""
        value = (filename or "").strip()
        if not value:
            return "generated_document.docx"

        stem = value
        ext = "docx"
        if "." in value:
            left, right = value.rsplit(".", 1)
            stem = left or stem
            candidate_ext = re.sub(r"[^A-Za-z0-9]", "", right.lower())
            if candidate_ext:
                ext = candidate_ext

        safe_stem = re.sub(r"[^A-Za-z0-9_-]", "_", stem)
        safe_stem = re.sub(r"_+", "_", safe_stem).strip("_-")
        if not safe_stem:
            safe_stem = "generated_document"

        return f"{safe_stem}.{ext}"

    @staticmethod
    def _generate_docx_bytes(request: FileGenerationRequest) -> bytes:
        try:
            document = DocxDocument()
            document.add_heading(request.filename, level=1)
            document.add_paragraph(request.content)

            buffer = BytesIO()
            document.save(buffer)
            return buffer.getvalue()
        except Exception as exc:
            raise FileGenerationError(
                message="Failed to generate DOCX artifact",
                details={"filename": request.filename, "error": str(exc)},
            ) from exc

    async def generate_docx(self, request: FileGenerationRequest) -> GeneratedFile:
        docx_bytes = self._generate_docx_bytes(request)
        safe_filename = self._sanitize_filename(request.filename)

        upload_url_endpoint = self._build_upload_url_endpoint()
        upload_request_payload: Dict[str, Any] = {
            "fileName": safe_filename,
            "mimeType": request.content_type,
            "bucket": self._resolve_bucket(request),
        }

        try:
            async with self._build_client() as client:
                upload_url_response = await client.post(
                    upload_url_endpoint,
                    headers=self._build_json_headers(),
                    json=upload_request_payload,
                )

                if upload_url_response.status_code >= 400:
                    raise FileGenerationError(
                        message="Files service upload-url request failed",
                        details={
                            "status_code": upload_url_response.status_code,
                            "url": upload_url_endpoint,
                            "response": upload_url_response.text,
                        },
                    )

                upload_url_payload = upload_url_response.json()
                if not isinstance(upload_url_payload, dict):
                    raise FileGenerationError(
                        message="Files service upload-url response is invalid",
                        details={"url": upload_url_endpoint},
                    )

                file_id = str(upload_url_payload.get("fileId") or "").strip()
                upload_url = str(upload_url_payload.get("uploadUrl") or "").strip()
                http_method = (
                    str(upload_url_payload.get("httpMethod") or "PUT").strip().upper()
                )

                if not file_id or not upload_url:
                    raise FileGenerationError(
                        message="Files service upload-url response missing required fields",
                        details={"url": upload_url_endpoint},
                    )

                if http_method not in {"PUT", "POST"}:
                    raise FileGenerationError(
                        message="Unsupported upload HTTP method from Files service",
                        details={"http_method": http_method},
                    )

                upload_response = await client.request(
                    method=http_method,
                    url=upload_url,
                    headers={"Content-Type": request.content_type},
                    content=docx_bytes,
                )

                if upload_response.status_code >= 400:
                    raise FileGenerationError(
                        message="Upload to blob storage failed",
                        details={
                            "file_id": file_id,
                            "status_code": upload_response.status_code,
                            "http_method": http_method,
                        },
                    )

                return GeneratedFile(
                    file_id=file_id,
                    filename=safe_filename,
                    content_type=request.content_type,
                    size_bytes=len(docx_bytes),
                )
        except FileGenerationError:
            raise
        except httpx.TimeoutException as exc:
            raise FileGenerationError(
                message="Files service request timed out",
                details={"url": upload_url_endpoint},
            ) from exc
        except httpx.HTTPError as exc:
            raise FileGenerationError(
                message="Files service HTTP request failed",
                details={"url": upload_url_endpoint, "error": str(exc)},
            ) from exc
        except ValueError as exc:
            raise FileGenerationError(
                message="Failed to parse Files service response",
                details={"url": upload_url_endpoint},
            ) from exc

    async def health_check(self) -> bool:
        if not self._base_url:
            return False

        try:
            async with self._build_client() as client:
                response = await client.get(self._base_url)
            return response.status_code < 500
        except Exception:
            return False


# Backward-compatible alias (legacy imports may still reference this name).
MockFileGenerationService = HTTPFileGenerationService


file_generation_service = HTTPFileGenerationService()
